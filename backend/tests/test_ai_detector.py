"""Tests for AI Detector engine and export endpoints."""
import tempfile
from pathlib import Path
from fastapi.testclient import TestClient
from app.ai_detector import evaluate_ai_footprint, compare_ai_footprint
from app.job_store import JobStore
from app.main import create_app
from app.docx_exporter import markdown_to_docx


def test_evaluate_ai_footprint_distinguishes_styles():
    # Robotic AI style text
    ai_text = (
        "In today's fast-paced digital world, it is crucial to leverage cutting-edge technology. "
        "Furthermore, we must delve into the intricate tapestry of modern innovations. "
        "In contrast, traditional paradigms fail to foster sustainable solutions. "
        "Moreover, this serves as a testament to our sophisticated operational efficiency."
    )
    res_ai = evaluate_ai_footprint(ai_text)
    assert res_ai.ai_percentage > 60.0
    assert len(res_ai.cliches_found) >= 3

    # Conversational human style text
    human_text = (
        "Say someone buys a watch online and it never shows up. What happens? "
        "A standard chatbot just spits out a tracking link. Useful, sure, but limited. "
        "An AI agent does much more: it spots the courier delay, pings dispatch, and fixes things on the spot. "
        "Big difference. That's why teams are switching."
    )
    res_human = evaluate_ai_footprint(human_text)
    assert res_human.ai_percentage < 15.0
    assert res_human.contraction_count >= 1
    assert res_human.burstiness_score >= 5.0


def test_ai_detection_endpoint_and_docx_export():
    with tempfile.TemporaryDirectory() as directory:
        db = Path(directory) / 'jobs.db'
        store = JobStore(db)
        app = create_app(store)
        client = TestClient(app)

        # Create and complete a job
        row = store.create("This is the original text draft about technology.", "local", [])
        job_id = row['job_id']
        store.complete(
            job_id,
            {
                "humanized_text": "# High Impact Technology\n\n*By Tech Team | March 2026*\n\nHere's how things work in reality.",
                "rules_applied": [],
                "audit": {"word_count": 14, "target_min": 10, "target_max": 20, "review_flags": []},
                "infographics": []
            }
        )

        # Test AI Detection endpoint
        res = client.get(f'/api/jobs/{job_id}/ai-detection')
        assert res.status_code == 200
        data = res.json()
        assert 'original' in data
        assert 'humanized' in data
        assert 'ai_percentage' in data['original']
        assert 'ai_percentage' in data['humanized']

        # Test DOCX export endpoint
        docx_res = client.get(f'/api/jobs/{job_id}/export/docx')
        assert docx_res.status_code == 200
        assert 'openxmlformats' in docx_res.headers['content-type']
        assert len(docx_res.content) > 500

        # Test Markdown export endpoint
        md_res = client.get(f'/api/jobs/{job_id}/export/markdown')
        assert md_res.status_code == 200
        assert 'High Impact Technology' in md_res.text


def test_markdown_to_docx_rich_formatting():
    import io
    import docx

    md_input = """# AI Agents vs Chatbots: Key Differences
*By Editorial Team | Updated: March 2026 | 6 min read*

Table of Contents
1. • [Overview](#overview)
2. • [Key Distinctions](#key-distinctions)

## Key Distinctions

Here is the breakdown:

| Feature | Chatbots | AI Agents |
|---|---|---|
| Primary Role | Answer FAQs | Complete tasks |
| Autonomy | Low | High |

> 📊 **Infographic / Visual Anchor:** [Summary of process flow comparison]

- Bullet item one with **bold text**
- Bullet item two with *italic text* and [link](https://example.com)
"""
    docx_bytes = markdown_to_docx(md_input)
    assert len(docx_bytes) > 2000

    doc = docx.Document(io.BytesIO(docx_bytes))
    # Assert real table generated
    assert len(doc.tables) >= 2  # data table + callout table
    
    # Check data table
    data_table = doc.tables[0]
    assert len(data_table.rows) == 3
    assert [c.text.strip() for c in data_table.rows[0].cells] == ['Feature', 'Chatbots', 'AI Agents']
    assert [c.text.strip() for c in data_table.rows[1].cells] == ['Primary Role', 'Answer FAQs', 'Complete tasks']
    
    # Check paragraphs for clean text without markdown leakage
    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "AI Agents vs Chatbots: Key Differences" in all_text
    assert "By Editorial Team | Updated: March 2026 | 6 min read" in all_text
    assert "#" not in all_text
    assert "**" not in all_text
    assert "[Overview](#overview)" not in all_text
    assert "Overview" in all_text

